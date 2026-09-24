"""
Convertidor de Markdown a DOCX sobrio y limpio (solo texto).
Lattice Bio.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

MD_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "PRESENTACION_CLINICA_ODONTOLOGIA_LATTICE_BIO.md")
DOCX_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "PRESENTACION_CLINICA_ODONTOLOGIA_LATTICE_BIO.docx")


def build_clean_docx():
    doc = Document()

    # Márgenes estándar (2.5 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Estilo de texto general
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        stripped = line.strip()
        if not stripped or stripped == '---':
            continue

        # Título principal #
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)

        # Subtítulos ##
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)

        # Sub-encabezados ###
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)

        # Viñetas -
        elif stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            text_bullet = stripped[2:].replace('**', '')
            p.add_run(text_bullet)

        # Texto normal
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            clean_text = stripped.replace('**', '')
            p.add_run(clean_text)

    doc.save(DOCX_PATH)
    print(f"[OK] Documento limpio DOCX generado en: {DOCX_PATH}")


if __name__ == '__main__':
    build_clean_docx()
